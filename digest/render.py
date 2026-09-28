import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .util import split_sentences


def _set_cjk(run, font: str, size: float, bold: bool = False, color: tuple[int, int, int] | None = None) -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _first_line_indent(paragraph, chars: int = 200) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.set(qn("w:firstLine"), "480")


def paragraphs_from_text(body: str, sentences_per_paragraph: int = 3) -> list[str]:
    flat = re.sub(r"\s+", " ", body.replace("\n", " ")).strip()
    sentences = split_sentences(flat)
    out: list[str] = []
    cur: list[str] = []
    for s in sentences:
        cur.append(s)
        if len(cur) >= sentences_per_paragraph:
            out.append(" ".join(cur))
            cur = []
    if cur:
        out.append(" ".join(cur))
    return out


def build_docx(path: Path, meta: dict, paragraphs: list[str]) -> Path:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cjk(head.add_run(meta["title_zh"]), "黑体", 18, bold=True)

    line1 = doc.add_paragraph()
    line1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cjk(line1.add_run(meta["meta_line"]), "宋体", 10.5, color=(0x59, 0x59, 0x59))

    if meta.get("orig_title"):
        line2 = doc.add_paragraph()
        line2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cjk(line2.add_run(meta["orig_title"]), "宋体", 9.5, color=(0x80, 0x80, 0x80))

    blank = doc.add_paragraph()
    blank.paragraph_format.space_after = Pt(2)

    for text in paragraphs:
        p = doc.add_paragraph()
        run = p.add_run(text)
        _set_cjk(run, "宋体", 12)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        _first_line_indent(p)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(14)
    _set_cjk(note.add_run(meta["note"]), "宋体", 9, color=(0x80, 0x80, 0x80))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
